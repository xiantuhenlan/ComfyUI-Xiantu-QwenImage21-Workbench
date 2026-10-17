from __future__ import annotations

import asyncio
import csv
import ctypes
import hashlib
import gc
import ipaddress
import io
import json
import mimetypes
import os
import re
import shutil
import subprocess
import tempfile
import time
import uuid
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path

from aiohttp import web

import folder_paths
from server import PromptServer


_IMAGE_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".bmp",
    ".gif",
    ".tif",
    ".tiff",
}
_IMPORT_SUBFOLDER = "xiantu_qwen_image21"
_VIDEO_EDITOR_SUBFOLDER = "xiantu_video_editor"
_MEDIA_EXTENSIONS = _IMAGE_EXTENSIONS | {
    ".mp4", ".mov", ".mkv", ".webm", ".avi", ".m4v",
    ".wav", ".mp3", ".flac", ".m4a", ".ogg", ".aac", ".wma",
}
_ASSET_ROOT_NAME = "闲兔资产库"
_ASSET_METADATA = "metadata.json"
_ASSET_LOCK = "metadata.lock"
_ASSET_CATEGORIES = {
    "character": "角色资产",
    "scene": "场景资产",
    "prop": "道具资产",
    "other": "其他资产",
    "camera": "运镜资产",
    "blocking": "站位资产",
}
_REPARSE_POINT = 0x400
_MAX_COMMON_ASSET_BYTES = 128 * 1024 * 1024
_MAX_SHEET_BYTES = 32 * 1024 * 1024


def _sheet_cell_text(value, number_format=""):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime,)):
        return value.isoformat(sep=" ")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        zero_format = str(number_format or "").strip()
        if isinstance(value, int) and re.fullmatch(r"0+", zero_format):
            return f"{value:0{len(zero_format)}d}"
        return f"{value:g}" if isinstance(value, float) else str(value)
    return str(value)


def _select_sheet_rows(candidates):
    first_nonempty = None
    for name, rows in candidates:
        useful = [row for row in rows if any(str(value).strip() for value in row)]
        if not useful:
            continue
        if first_nonempty is None:
            first_nonempty = (name, useful)
        for index, row in enumerate(useful[:20]):
            headers = {re.sub(r"\s+", "", str(value or "")) for value in row}
            if "提示词" in headers and ({"分镜编号", "序号", "编号", "镜头编号"} & headers):
                return name, useful[index:]
    if first_nonempty:
        return first_nonempty
    raise ValueError("表格没有可读取的数据")


def _parse_xlsx_rows(data):
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("缺少 openpyxl，无法读取 .xlsx 表格") from error
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    candidates = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows():
                rows.append([_sheet_cell_text(cell.value, cell.number_format) for cell in row[:128]])
                if len(rows) >= 10000:
                    break
            candidates.append((sheet.title, rows))
    finally:
        workbook.close()
    return _select_sheet_rows(candidates)


def _parse_xls_rows(data):
    try:
        import xlrd
    except ImportError as error:
        raise RuntimeError("缺少 xlrd，无法读取 .xls 表格；请重新安装插件依赖") from error
    workbook = xlrd.open_workbook(file_contents=data, formatting_info=True)
    candidates = []
    for sheet in workbook.sheets():
        rows = []
        for row_index in range(min(sheet.nrows, 10000)):
            values = []
            for column_index in range(min(sheet.ncols, 128)):
                cell = sheet.cell(row_index, column_index)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, workbook.datemode)
                number_format = ""
                try:
                    format_key = workbook.xf_list[cell.xf_index].format_key
                    number_format = workbook.format_map[format_key].format_str
                except (AttributeError, IndexError, KeyError):
                    pass
                values.append(_sheet_cell_text(value, number_format))
            rows.append(values)
        candidates.append((sheet.name, rows))
    return _select_sheet_rows(candidates)


def _parse_csv_rows(data):
    text = None
    for encoding in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("CSV 编码无法识别，请保存为 UTF-8、GBK 或 UTF-16")
    rows = [[str(value) for value in row[:128]] for row in csv.reader(io.StringIO(text))]
    return _select_sheet_rows([("CSV", rows[:10000])])


class _AssetLibraryError(RuntimeError):
    def __init__(self, code, message, status=400):
        super().__init__(message)
        self.code = code
        self.status = status


def _asset_error(error):
    if isinstance(error, _AssetLibraryError):
        return web.json_response({"error": str(error), "code": error.code}, status=error.status)
    return web.json_response({"error": str(error), "code": "asset_operation_failed"}, status=500)


def _is_reparse(path: Path) -> bool:
    try:
        return bool(getattr(os.lstat(path), "st_file_attributes", 0) & _REPARSE_POINT)
    except OSError:
        return False


def _asset_root() -> Path:
    configured = str(os.environ.get("XIANTU_ASSET_LIBRARY_ROOT") or "").strip()
    candidates = [Path(configured)] if configured else []
    comfy_root = Path(folder_paths.base_path).resolve()
    candidates.extend(parent / _ASSET_ROOT_NAME for parent in (comfy_root, *comfy_root.parents))
    for candidate in candidates:
        try:
            resolved = candidate.resolve()
        except OSError:
            continue
        if resolved.is_dir() and not resolved.is_symlink() and not _is_reparse(resolved):
            return resolved
    raise _AssetLibraryError(
        "asset_root_not_found",
        f"公共资产库未建立或未定位：请在启动器包根建立“{_ASSET_ROOT_NAME}”文件夹",
        404,
    )


def _utc_now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")


def _asset_relative(category, filename):
    return f"{_ASSET_CATEGORIES[category]}/{filename}".replace("\\", "/")


def _checked_asset_path(root: Path, relative_path: str, must_exist=True) -> Path:
    normalized = str(relative_path or "").replace("\\", "/")
    parts = Path(normalized).parts
    if len(parts) != 2 or parts[0] not in _ASSET_CATEGORIES.values() or any(part in {"", ".", ".."} for part in parts):
        raise _AssetLibraryError("asset_path_outside_root", "资产相对路径无效")
    parent = (root / parts[0]).resolve()
    try:
        parent.relative_to(root)
    except ValueError as error:
        raise _AssetLibraryError("asset_path_outside_root", "资产路径越过公共资产库") from error
    if not parent.is_dir() or parent.is_symlink() or _is_reparse(parent):
        raise _AssetLibraryError("asset_reparse_point_rejected", "资产分类目录不是普通目录")
    candidate = parent / parts[1]
    if must_exist:
        if not candidate.is_file():
            raise _AssetLibraryError("asset_source_missing", f"公共资产不存在：{normalized}", 404)
        if candidate.is_symlink() or _is_reparse(candidate):
            raise _AssetLibraryError("asset_reparse_point_rejected", "拒绝读取链接或重解析点资产")
    return candidate


def _sha256_file(path: Path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def _mime_type(path: Path):
    return mimetypes.guess_type(path.name)[0] or "application/octet-stream"


def _read_asset_index(root: Path):
    path = root / _ASSET_METADATA
    if not path.is_file():
        return {"schemaVersion": 1, "assets": []}, True, "未找到 metadata.json，已根据六类文件夹重建"
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(payload, dict) or not isinstance(payload.get("assets"), list):
            raise ValueError("metadata.json 结构无效")
        return {"schemaVersion": 1, "assets": payload["assets"]}, False, ""
    except (OSError, ValueError, json.JSONDecodeError) as error:
        backup = root / f"metadata.corrupt-{int(time.time())}.json"
        try:
            shutil.copy2(path, backup)
        except OSError:
            pass
        return {"schemaVersion": 1, "assets": []}, True, f"metadata.json 已损坏，已保留副本并重建：{error}"


def _write_asset_index(root: Path, assets):
    destination = root / _ASSET_METADATA
    temporary = root / f"metadata.json.tmp-{os.getpid()}-{uuid.uuid4().hex}"
    payload = json.dumps({"schemaVersion": 1, "assets": assets}, ensure_ascii=False, indent=2).encode("utf-8")
    try:
        with temporary.open("xb") as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
        if os.name == "nt":
            move_file = ctypes.windll.kernel32.MoveFileExW
            move_file.argtypes = [ctypes.c_wchar_p, ctypes.c_wchar_p, ctypes.c_uint32]
            move_file.restype = ctypes.c_int
            result = move_file(str(temporary), str(destination), 0x1 | 0x8)
            if not result:
                raise ctypes.WinError()
        else:
            os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)


@contextmanager
def _asset_metadata_lock(root: Path):
    lock_path = root / _ASSET_LOCK
    for _ in range(100):
        try:
            descriptor = os.open(lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
                handle.write(f"{os.getpid()}:{int(time.time())}\n")
            break
        except FileExistsError:
            try:
                if time.time() - lock_path.stat().st_mtime > 120:
                    lock_path.unlink(missing_ok=True)
            except OSError:
                pass
            time.sleep(.02)
    else:
        raise _AssetLibraryError("asset_metadata_locked", "资产库正在被其他窗口修改，请稍后重试", 409)
    try:
        yield
    finally:
        lock_path.unlink(missing_ok=True)


def _reconcile_asset_library(root: Path):
    index, recovered, recovery_message = _read_asset_index(root)
    indexed_by_path = {}
    indexed_by_hash = {}
    for item in index["assets"]:
        if not isinstance(item, dict):
            continue
        relative = str(item.get("relativePath") or "").replace("\\", "/")
        indexed_by_path[relative] = item
        content_hash = str(item.get("contentHash") or "").upper()
        if content_hash:
            indexed_by_hash.setdefault(content_hash, []).append(item)
    used_ids = set()
    assets = []
    for category, directory in _ASSET_CATEGORIES.items():
        category_root = root / directory
        if not category_root.is_dir() or category_root.is_symlink() or _is_reparse(category_root):
            continue
        for path in category_root.iterdir():
            if not path.is_file() or path.is_symlink() or _is_reparse(path):
                continue
            relative = _asset_relative(category, path.name)
            content_hash = _sha256_file(path)
            prior = indexed_by_path.get(relative)
            if prior is None:
                prior = next((item for item in indexed_by_hash.get(content_hash, []) if str(item.get("id")) not in used_ids), None)
            now = _utc_now()
            asset_id = str(prior.get("id") if prior else "").strip() or f"asset-{uuid.uuid4().hex[:24]}"
            used_ids.add(asset_id)
            name = str(prior.get("name") if prior else "").strip() or path.stem or "未命名资产"
            item = {
                "id": asset_id,
                "relativePath": relative,
                "name": name,
                "note": str(prior.get("note") if prior else "").strip(),
                "category": category,
                "mimeType": _mime_type(path),
                "contentHash": content_hash,
                "createdAt": str(prior.get("createdAt") if prior else "").strip() or now,
                "updatedAt": str(prior.get("updatedAt") if prior else "").strip() or now,
                "sizeBytes": path.stat().st_size,
            }
            assets.append(item)
    assets.sort(key=lambda item: (item["updatedAt"], item["name"]), reverse=True)
    _write_asset_index(root, [{key: value for key, value in item.items() if key != "sizeBytes"} for item in assets])
    return {
        "schemaVersion": 1,
        "root": str(root),
        "assets": assets,
        "recoveredMetadata": recovered,
        "recoveryMessage": recovery_message,
    }


def _asset_snapshot():
    root = _asset_root()
    with _asset_metadata_lock(root):
        return root, _reconcile_asset_library(root)


def _safe_file_name(value):
    value = Path(str(value or "").strip()).name
    if not value or value in {".", ".."} or any(char in value for char in '<>:"/\\|?*') or value.endswith((" ", ".")):
        raise _AssetLibraryError("asset_invalid_filename", "资产文件名无效")
    return value


def _unique_asset_path(parent: Path, requested):
    requested = Path(_safe_file_name(requested))
    for index in range(10000):
        name = requested.name if index == 0 else f"{requested.stem} ({index}){requested.suffix}"
        candidate = parent / name
        if not candidate.exists():
            return candidate
    raise _AssetLibraryError("asset_duplicate_limit", "同名资产过多，无法保存")


def _snapshot_asset(snapshot, asset_id):
    asset = next((item for item in snapshot["assets"] if item["id"] == str(asset_id)), None)
    if asset is None:
        raise _AssetLibraryError("asset_source_missing", "没有找到公共资产", 404)
    return asset


@PromptServer.instance.routes.get("/xiantu/video-editor/common-assets")
async def list_common_assets(_request):
    try:
        _root, snapshot = _asset_snapshot()
        return web.json_response(snapshot)
    except Exception as error:
        return _asset_error(error)


@PromptServer.instance.routes.get("/xiantu/video-editor/common-assets/content")
async def common_asset_content(request):
    try:
        root = _asset_root()
        index, _recovered, _message = _read_asset_index(root)
        asset = next((item for item in index["assets"] if isinstance(item, dict) and str(item.get("id")) == str(request.query.get("id") or "")), None)
        if asset is None:
            raise _AssetLibraryError("asset_source_missing", "没有找到公共资产", 404)
        path = _checked_asset_path(root, asset.get("relativePath"))
        return web.FileResponse(path, headers={"Cache-Control": "no-cache"})
    except Exception as error:
        return _asset_error(error)


@PromptServer.instance.routes.post("/xiantu/video-editor/common-assets/import")
async def import_common_asset(request):
    try:
        reader = await request.multipart()
        fields = {"category": "other", "name": "", "note": ""}
        file_name = ""
        file_bytes = None
        while field := await reader.next():
            if field.name == "media" and field.filename:
                file_name = Path(field.filename).name
                file_bytes = await field.read(decode=False)
            elif field.name in fields:
                fields[field.name] = await field.text()
        if not file_bytes or len(file_bytes) > _MAX_COMMON_ASSET_BYTES:
            raise _AssetLibraryError("asset_import_invalid", "资产为空或超过 128 MiB")
        category = fields["category"] if fields["category"] in _ASSET_CATEGORIES else "other"
        root = _asset_root()
        with _asset_metadata_lock(root):
            snapshot = _reconcile_asset_library(root)
            destination = _unique_asset_path(root / _ASSET_CATEGORIES[category], file_name)
            with destination.open("xb") as handle:
                handle.write(file_bytes)
                handle.flush()
                os.fsync(handle.fileno())
            snapshot = _reconcile_asset_library(root)
            relative = _asset_relative(category, destination.name)
            asset = next(item for item in snapshot["assets"] if item["relativePath"] == relative)
            asset["name"] = str(fields["name"] or "").strip()[:200] or destination.stem
            asset["note"] = str(fields["note"] or "").strip()[:2000]
            asset["updatedAt"] = _utc_now()
            _write_asset_index(root, [{key: value for key, value in item.items() if key != "sizeBytes"} for item in snapshot["assets"]])
        return web.json_response({"ok": True, "asset": asset})
    except Exception as error:
        return _asset_error(error)


@PromptServer.instance.routes.post("/xiantu/video-editor/common-assets/update")
async def update_common_asset(request):
    try:
        payload = await request.json()
        root = _asset_root()
        with _asset_metadata_lock(root):
            snapshot = _reconcile_asset_library(root)
            asset = _snapshot_asset(snapshot, payload.get("id"))
            new_category = str(payload.get("category") or asset["category"])
            if new_category not in _ASSET_CATEGORIES:
                raise _AssetLibraryError("asset_category_invalid", "未知资产分类")
            if new_category != asset["category"]:
                source = _checked_asset_path(root, asset["relativePath"])
                destination = _unique_asset_path(root / _ASSET_CATEGORIES[new_category], source.name)
                os.replace(source, destination)
                asset["relativePath"] = _asset_relative(new_category, destination.name)
                asset["category"] = new_category
            if "name" in payload:
                name = str(payload.get("name") or "").strip()
                if not name or len(name) > 200 or "\n" in name or "\r" in name:
                    raise _AssetLibraryError("asset_name_invalid", "资产名称必须为 1 至 200 个字符且不能换行")
                asset["name"] = name
            if "note" in payload:
                note = str(payload.get("note") or "").strip()
                if len(note) > 2000:
                    raise _AssetLibraryError("asset_note_invalid", "资产备注不能超过 2000 个字符")
                asset["note"] = note
            asset["updatedAt"] = _utc_now()
            _write_asset_index(root, [{key: value for key, value in item.items() if key != "sizeBytes"} for item in snapshot["assets"]])
        return web.json_response({"ok": True, "asset": asset})
    except Exception as error:
        return _asset_error(error)


@PromptServer.instance.routes.post("/xiantu/video-editor/common-assets/delete")
async def delete_common_assets(request):
    try:
        payload = await request.json()
        ids = {str(value) for value in payload.get("ids", [])}
        root = _asset_root()
        with _asset_metadata_lock(root):
            snapshot = _reconcile_asset_library(root)
            for asset in snapshot["assets"]:
                if asset["id"] in ids:
                    _checked_asset_path(root, asset["relativePath"]).unlink()
            snapshot = _reconcile_asset_library(root)
        return web.json_response({"ok": True, "snapshot": snapshot})
    except Exception as error:
        return _asset_error(error)


@PromptServer.instance.routes.post("/xiantu/video-editor/common-assets/apply")
async def apply_common_assets(request):
    """The only common-library operation allowed to write into ComfyUI/input."""
    try:
        payload = await request.json()
        ids = [str(value) for value in payload.get("ids", [])]
        root, snapshot = _asset_snapshot()
        input_root = Path(folder_paths.get_input_directory()).resolve()
        target_dir = input_root / _VIDEO_EDITOR_SUBFOLDER
        target_dir.mkdir(parents=True, exist_ok=True)
        applied = []
        for asset_id in ids:
            asset = _snapshot_asset(snapshot, asset_id)
            source = _checked_asset_path(root, asset["relativePath"])
            safe_stem = "".join(char if char.isalnum() or char in "-_" else "_" for char in source.stem).strip("_") or "asset"
            target_name = f"{safe_stem}_{asset['contentHash'][:12].lower()}{source.suffix.lower()}"
            target = target_dir / target_name
            if not target.exists():
                shutil.copy2(source, target)
            applied.append({
                "id": asset["id"],
                "name": asset["name"],
                "note": asset["note"],
                "category": asset["category"],
                "contentHash": asset["contentHash"],
                "sourceRelativePath": asset["relativePath"],
                "filename": target_name,
                "subfolder": _VIDEO_EDITOR_SUBFOLDER,
                "type": "input",
                "mimeType": asset["mimeType"],
            })
        return web.json_response({"ok": True, "assets": applied})
    except OSError as error:
        return _asset_error(_AssetLibraryError("asset_input_copy_failed", f"复制资产到 ComfyUI/input 失败：{error}", 500))
    except Exception as error:
        return _asset_error(error)


@PromptServer.instance.routes.post("/xiantu/video-editor/parse-sheet")
async def parse_video_editor_sheet(request):
    try:
        reader = await request.multipart()
        field = await reader.next()
    except Exception:
        return _error("上传内容不是有效的表格文件")
    if field is None or field.name != "sheet" or not field.filename:
        return _error("没有收到表格文件")

    filename = Path(field.filename).name
    suffix = Path(filename).suffix.lower()
    if suffix not in {".csv", ".xls", ".xlsx"}:
        return _error("仅支持 CSV、XLS 和 XLSX 表格")
    content = bytearray()
    while True:
        chunk = await field.read_chunk(size=1024 * 1024)
        if not chunk:
            break
        content.extend(chunk)
        if len(content) > _MAX_SHEET_BYTES:
            return _error("表格文件不能超过 32 MB", status=413)
    try:
        if suffix == ".xlsx":
            sheet_name, rows = _parse_xlsx_rows(bytes(content))
        elif suffix == ".xls":
            sheet_name, rows = _parse_xls_rows(bytes(content))
        else:
            sheet_name, rows = _parse_csv_rows(bytes(content))
        return web.json_response({"filename": filename, "sheet": sheet_name, "rows": rows})
    except Exception as error:
        return _error(f"读取表格失败：{error}")


@PromptServer.instance.routes.post("/xiantu/video-editor/upload-media")
async def upload_video_editor_media(request):
    """Persist a browser-selected media file under ComfyUI/input.

    Browser object URLs only live until the workbench is closed.  Keeping the
    original upload in the input directory gives both the reopened workbench
    and the Python H3 encoder a stable source of truth.
    """
    try:
        reader = await request.multipart()
        field = await reader.next()
    except Exception:
        return _error("上传内容不是有效的文件表单")
    if field is None or field.name != "media" or not field.filename:
        return _error("没有收到素材文件")

    original_name = Path(field.filename).name
    suffix = Path(original_name).suffix.lower()
    if suffix not in _MEDIA_EXTENSIONS:
        return _error(f"不支持的素材格式：{suffix or '无扩展名'}")

    input_root = Path(folder_paths.get_input_directory()).resolve()
    target_dir = input_root / _VIDEO_EDITOR_SUBFOLDER
    target_dir.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256()
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(dir=target_dir, suffix=suffix, delete=False) as handle:
            temp_path = Path(handle.name)
            while True:
                chunk = await field.read_chunk(size=1024 * 1024)
                if not chunk:
                    break
                digest.update(chunk)
                handle.write(chunk)
        safe_stem = "".join(char if char.isalnum() or char in "-_" else "_" for char in Path(original_name).stem).strip("_") or "media"
        target_name = f"{safe_stem}_{digest.hexdigest()[:12]}{suffix}"
        target = target_dir / target_name
        if target.exists():
            temp_path.unlink(missing_ok=True)
        else:
            os.replace(temp_path, target)
    except OSError as error:
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)
        return _error(f"保存上传素材失败：{error}", status=500)

    mime_type = field.headers.get("Content-Type") or mimetypes.guess_type(original_name)[0] or "application/octet-stream"
    return web.json_response({
        "name": target_name,
        "subfolder": _VIDEO_EDITOR_SUBFOLDER,
        "type": "input",
        "mimeType": mime_type,
        "displayName": original_name,
    })


def _error(message: str, status: int = 400) -> web.Response:
    return web.json_response({"error": message}, status=status)


def _ffmpeg_executable():
    executable = shutil.which("ffmpeg")
    if executable:
        return executable
    try:
        import imageio_ffmpeg
        return imageio_ffmpeg.get_ffmpeg_exe()
    except Exception:
        return None


def _output_record_path(record):
    if not isinstance(record, dict):
        raise ValueError("输出视频记录无效")
    output_root = Path(folder_paths.get_output_directory()).resolve()
    name = Path(str(record.get("filename") or "")).name
    subfolder = str(record.get("subfolder") or "").replace("\\", "/").strip("/")
    candidate = (output_root / subfolder / name).resolve()
    candidate.relative_to(output_root)
    if not candidate.is_file():
        raise FileNotFoundError(name)
    return candidate


def _sync_finalized_name_to_history(old_name: str, new_name: str, subfolder: str) -> None:
    """Keep ComfyUI history/native previews aligned with an atomic rename."""
    queue = getattr(PromptServer.instance, "prompt_queue", None)
    history = getattr(queue, "history", None)
    mutex = getattr(queue, "mutex", None)
    if not isinstance(history, dict) or mutex is None:
        return

    normalized_subfolder = str(subfolder or "").replace("\\", "/").strip("/")

    def visit(value):
        if isinstance(value, dict):
            current_subfolder = str(value.get("subfolder") or "").replace("\\", "/").strip("/")
            if str(value.get("filename") or "") == old_name and current_subfolder == normalized_subfolder:
                value["filename"] = new_name
                if str(value.get("name") or "") == old_name:
                    value["name"] = new_name
            for child in value.values():
                visit(child)
        elif isinstance(value, list):
            for child in value:
                visit(child)

    with mutex:
        visit(history)


@PromptServer.instance.routes.post("/xiantu/video-editor/finalize-video-name")
async def finalize_video_name(request):
    """Replace VHS' fixed ``-audio`` suffix after muxing completes."""
    try:
        record = await request.json()
        output_root = Path(folder_paths.get_output_directory()).resolve()
        name = Path(str(record.get("filename") or "")).name
        subfolder = str(record.get("subfolder") or "").replace("\\", "/").strip("/")
        source = (output_root / subfolder / name).resolve()
        source.relative_to(output_root)
    except (ValueError, OSError) as error:
        return _error(f"最终视频路径无效：{error}")
    has_audio_suffix = source.stem.lower().endswith("-audio")
    muxed_stem = source.stem[:-6] if has_audio_suffix else source.stem
    # VHS appends a five-digit counter and may also emit a silent MP4 and a
    # metadata PNG.  The director promises one authoritative video per shot,
    # so publish the muxed result under the stable prefix and remove only the
    # matching VHS intermediates after the atomic replacement succeeds.
    stable_stem = re.sub(r"_\d{5}$", "", muxed_stem)
    target = source.with_name(f"{stable_stem}{source.suffix}")
    # This route is deliberately idempotent. Saved projects may still contain
    # VHS' old -audio name after the muxed file has already been finalized.
    if not source.is_file() and target.is_file():
        if target.stat().st_size <= 0:
            return _error("视频仍在生成，请完成后再播放", status=409)
        _sync_finalized_name_to_history(source.name, target.name, subfolder)
        return web.json_response({
            "ok": True,
            "filename": target.name,
            "subfolder": subfolder,
            "type": "output",
            "size": target.stat().st_size,
        })
    if not source.is_file():
        return _error(f"找不到最终视频：{source.name}", status=404)
    if source == target:
        if source.stat().st_size <= 0:
            return _error("视频仍在生成，请完成后再播放", status=409)
        return web.json_response({
            "ok": True,
            "filename": source.name,
            "subfolder": subfolder,
            "type": "output",
            "size": source.stat().st_size,
        })
    deadline = time.monotonic() + 12.0
    while True:
        try:
            # VHS first creates this same-name silent intermediate. It is not
            # the deliverable; the muxed -audio file is the final video. Keep
            # the old target in place until the atomic replacement succeeds;
            # never delete it first when Windows still has either file open.
            os.replace(source, target)
            break
        except PermissionError as error:
            if time.monotonic() >= deadline:
                return _error(f"视频编码器仍在占用文件，稍后将再次整理文件名：{error}", status=409)
            await asyncio.sleep(0.25)
        except OSError as error:
            return _error(f"最终视频改名失败：{error}", status=500)
    for sibling in target.parent.glob(f"{stable_stem}_?????*"):
        if sibling == target or not sibling.is_file():
            continue
        if sibling.suffix.lower() not in {".mp4", ".webm", ".mov", ".mkv", ".avi", ".m4v", ".png"}:
            continue
        try:
            sibling.unlink()
        except OSError:
            pass
    _sync_finalized_name_to_history(source.name, target.name, subfolder)
    return web.json_response({
        "ok": True,
        "filename": target.name,
        "subfolder": subfolder,
        "type": "output",
        "size": target.stat().st_size,
    })


@PromptServer.instance.routes.post("/xiantu/video-editor/rename-video")
async def rename_video(request):
    """Rename an output video and return the authoritative output record."""
    try:
        record = await request.json()
        source = _output_record_path(record)
    except FileNotFoundError as error:
        return _error(f"找不到待重命名视频：{error}", status=404)
    except (ValueError, OSError) as error:
        return _error(f"视频路径无效：{error}")
    requested = Path(str(record.get("new_name") or "").strip()).name
    if not requested or requested in {".", ".."}:
        return _error("请输入有效文件名")
    video_suffixes = {".mp4", ".webm", ".mov", ".mkv", ".avi", ".m4v"}
    if Path(requested).suffix.lower() not in video_suffixes:
        requested = f"{requested}{source.suffix}"
    target = source.with_name(requested)
    if target == source:
        return web.json_response({
            "ok": True,
            "filename": source.name,
            "subfolder": str(record.get("subfolder") or "").replace("\\", "/").strip("/"),
            "type": "output",
        })
    if target.exists():
        return _error(f"同名文件已存在：{target.name}", status=409)
    try:
        os.replace(source, target)
    except OSError as error:
        return _error(f"视频重命名失败：{error}", status=500)
    return web.json_response({
        "ok": True,
        "filename": target.name,
        "subfolder": str(record.get("subfolder") or "").replace("\\", "/").strip("/"),
        "type": "output",
    })


@PromptServer.instance.routes.get("/xiantu/video-editor/face-repair-models")
async def face_repair_models(_request):
    """Report the real on-disk state of every face-repair model."""
    models_root = Path(getattr(folder_paths, "models_dir", Path(folder_paths.base_path) / "models")).resolve()
    expected = {
        "faceRestore": models_root / "facerestore" / "GFPGANv1.4.pth",
        "faceDetector": models_root / "ultralytics" / "bbox" / "face_yolov8m.pt",
        "personDetector": models_root / "ultralytics" / "segm" / "person_yolov8m-seg.pt",
    }
    return web.json_response({
        key: {"installed": path.is_file() and path.stat().st_size > 0, "size": path.stat().st_size if path.is_file() else 0}
        for key, path in expected.items()
    })


@PromptServer.instance.routes.get("/xiantu/video-editor/hardware")
async def video_editor_hardware(_request):
    """Small local hardware snapshot for the director header."""
    data = {"cpu": None, "memory": None, "vram": None, "temperature": None}
    try:
        import psutil
        data["cpu"] = round(float(psutil.cpu_percent(interval=None)), 1)
        data["memory"] = round(float(psutil.virtual_memory().percent), 1)
    except Exception:
        pass
    try:
        completed = subprocess.run(
            ["nvidia-smi", "--query-gpu=memory.used,memory.total,temperature.gpu", "--format=csv,noheader,nounits"],
            capture_output=True, text=True, timeout=3, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0), check=False,
        )
        first = completed.stdout.strip().splitlines()[0]
        used, total, temperature = [float(item.strip()) for item in first.split(",")[:3]]
        data["vram"] = round(used / max(1.0, total) * 100.0, 1)
        data["temperature"] = round(temperature, 1)
    except Exception:
        pass
    return web.json_response(data)


@PromptServer.instance.routes.post("/xiantu/video-editor/cleanup")
async def video_editor_cleanup(request):
    """Release only the resources explicitly selected by the user."""
    try:
        payload = await request.json()
    except Exception:
        payload = {}
    released = []
    warnings = []
    if bool(payload.get("models")):
        try:
            import comfy.model_management as model_management
            model_management.unload_all_models()
            released.append("models")
        except Exception as error:
            warnings.append(f"卸载模型失败：{error}")
    if bool(payload.get("cache")):
        try:
            gc.collect()
            released.append("cache")
        except Exception as error:
            warnings.append(f"卸载缓存失败：{error}")
    if bool(payload.get("vram")):
        try:
            import torch
            if torch.cuda.is_available():
                torch.cuda.empty_cache()
                torch.cuda.ipc_collect()
            released.append("vram")
        except Exception as error:
            warnings.append(f"卸载显存失败：{error}")
    return web.json_response({"ok": True, "released": released, "warnings": warnings})


@PromptServer.instance.routes.post("/xiantu/video-editor/open-output")
async def video_editor_open_output(_request):
    try:
        path = Path(folder_paths.get_output_directory()).resolve()
        if os.name == "nt":
            os.startfile(path)  # type: ignore[attr-defined]
        else:
            subprocess.Popen(["xdg-open", str(path)])
        return web.json_response({"ok": True})
    except Exception as error:
        return _error(f"打开输出目录失败：{error}", status=500)


@PromptServer.instance.routes.post("/xiantu/video-editor/merge")
async def video_editor_merge(request):
    try:
        payload = await request.json()
        records = payload.get("videos") if isinstance(payload, dict) else None
    except Exception:
        records = None
    if not isinstance(records, list) or not records:
        return _error("请先选择要合并的视频")
    paths = []
    output_root = Path(folder_paths.get_output_directory()).resolve()
    for record in records:
        if not isinstance(record, dict):
            return _error("合并列表格式无效")
        name = Path(str(record.get("filename") or "")).name
        subfolder = str(record.get("subfolder") or "").replace("\\", "/").strip("/")
        candidate = (output_root / subfolder / name).resolve()
        try:
            candidate.relative_to(output_root)
        except ValueError:
            return _error("合并文件超出输出目录")
        if not candidate.is_file():
            return _error(f"找不到待合并视频：{name}", status=404)
        paths.append(candidate)
    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg:
        try:
            import imageio_ffmpeg
            ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
        except Exception:
            return _error("未找到 FFmpeg，无法合并视频", status=500)
    target_dir = output_root / "短剧-漫剧" / "合并成品"
    target_dir.mkdir(parents=True, exist_ok=True)
    index = 1
    while (target_dir / f"成品-{index:03d}.mp4").exists():
        index += 1
    target = target_dir / f"成品-{index:03d}.mp4"
    list_file = None
    try:
        with tempfile.NamedTemporaryFile("w", encoding="utf-8", suffix=".txt", delete=False) as handle:
            list_file = Path(handle.name)
            for path in paths:
                handle.write("file '" + str(path).replace("'", "'\\''") + "'\n")
        completed = subprocess.run(
            [ffmpeg, "-y", "-f", "concat", "-safe", "0", "-i", str(list_file), "-c", "copy", str(target)],
            capture_output=True, text=True, timeout=3600, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0), check=False,
        )
        if completed.returncode != 0:
            return _error(f"合并失败：{completed.stderr[-500:]}", status=500)
    finally:
        if list_file is not None:
            list_file.unlink(missing_ok=True)
    return web.json_response({
        "ok": True,
        "filename": target.name,
        "subfolder": "短剧-漫剧/合并成品",
        "type": "output",
    })


@PromptServer.instance.routes.post("/xiantu/qwen21/import-local-image")
async def import_local_image(request):
    if request.remote:
        try:
            is_loopback = ipaddress.ip_address(request.remote.split("%", 1)[0]).is_loopback
        except ValueError:
            is_loopback = False
        if not is_loopback:
            return _error("绝对路径图片只能在运行 ComfyUI 的本机导入", status=403)

    try:
        payload = await request.json()
    except Exception:
        return _error("请求内容不是有效的 JSON")

    raw_path = str(payload.get("path") or "").strip().strip('"')
    if not raw_path:
        return _error("图片绝对路径不能为空")

    source = Path(raw_path).expanduser()
    if not source.is_absolute():
        return _error("参考图必须填写绝对路径")

    try:
        source = source.resolve(strict=True)
    except (OSError, RuntimeError):
        return _error(f"找不到图片：{raw_path}", status=404)

    if not source.is_file():
        return _error(f"不是图片文件：{raw_path}")
    if source.suffix.lower() not in _IMAGE_EXTENSIONS:
        return _error(f"不支持的图片格式：{source.suffix or '无扩展名'}")

    input_root = Path(folder_paths.get_input_directory()).resolve()
    target_dir = input_root / _IMPORT_SUBFOLDER
    target_dir.mkdir(parents=True, exist_ok=True)

    path_key = str(source).casefold().encode("utf-8")
    digest = hashlib.sha256(path_key).hexdigest()[:12]
    target_name = f"{source.stem}_{digest}{source.suffix.lower()}"
    target = target_dir / target_name

    try:
        if source != target.resolve():
            shutil.copy2(source, target)
    except OSError as error:
        return _error(f"导入图片失败：{error}", status=500)

    return web.json_response(
        {
            "name": target_name,
            "subfolder": _IMPORT_SUBFOLDER,
            "type": "input",
            "displayName": source.name,
            "sourcePath": str(source),
        }
    )


@PromptServer.instance.routes.post("/xiantu/video-editor/import-local-media")
async def import_local_media(request):
    if request.remote:
        try:
            is_loopback = ipaddress.ip_address(request.remote.split("%", 1)[0]).is_loopback
        except ValueError:
            is_loopback = False
        if not is_loopback:
            return _error("绝对路径素材只能在运行 ComfyUI 的本机导入", status=403)

    try:
        payload = await request.json()
    except Exception:
        return _error("请求内容不是有效的 JSON")

    raw_path = str(payload.get("path") or "").strip().strip('"')
    if not raw_path:
        return _error("素材绝对路径不能为空")
    source = Path(raw_path).expanduser()
    if not source.is_absolute():
        return _error("素材必须填写绝对路径")
    try:
        source = source.resolve(strict=True)
    except (OSError, RuntimeError):
        return _error(f"找不到素材：{raw_path}", status=404)
    if not source.is_file():
        return _error(f"不是文件：{raw_path}")
    if source.suffix.lower() not in _MEDIA_EXTENSIONS:
        return _error(f"不支持的素材格式：{source.suffix or '无扩展名'}")

    input_root = Path(folder_paths.get_input_directory()).resolve()
    target_dir = input_root / _VIDEO_EDITOR_SUBFOLDER
    target_dir.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(str(source).casefold().encode("utf-8")).hexdigest()[:12]
    target_name = f"{source.stem}_{digest}{source.suffix.lower()}"
    target = target_dir / target_name
    try:
        if source != target.resolve():
            shutil.copy2(source, target)
    except OSError as error:
        return _error(f"导入素材失败：{error}", status=500)

    mime_type = mimetypes.guess_type(source.name)[0] or "application/octet-stream"
    return web.json_response({
        "name": target_name,
        "subfolder": _VIDEO_EDITOR_SUBFOLDER,
        "type": "input",
        "mimeType": mime_type,
        "displayName": source.name,
        "sourcePath": str(source),
    })
